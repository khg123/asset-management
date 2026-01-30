import os
import datetime
import random
import string
from io import BytesIO
from base64 import b64encode
# 设置环境变量以消除SQLAlchemy 2.0兼容性警告
os.environ['SQLALCHEMY_SILENCE_UBER_WARNING'] = '1'

from flask import Flask, render_template, redirect, url_for, request, flash, session
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, login_required, logout_user, current_user
from flask_wtf import FlaskForm
from wtforms import StringField, PasswordField, SubmitField, SelectField, IntegerField, DateField, FloatField, TextAreaField
from wtforms.validators import InputRequired, Length, EqualTo, Optional
from passlib.hash import sha256_crypt
from flask_bootstrap import Bootstrap
from dotenv import load_dotenv

# 加载环境变量
load_dotenv()

app = Flask(__name__)
app.config['SECRET_KEY'] = os.getenv('SECRET_KEY', 'dev-secret-key')
app.config['SQLALCHEMY_DATABASE_URI'] = os.getenv('DATABASE_URI', 'mysql+pymysql://root:password@localhost/asset_management')
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False
app.config['STATIC_FOLDER'] = 'static'

Bootstrap(app)
db = SQLAlchemy(app)
login_manager = LoginManager()
login_manager.init_app(app)
login_manager.login_view = 'login'

# 数据库模型
class User(UserMixin, db.Model):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(20), unique=True, nullable=False)
    password = db.Column(db.String(100), nullable=False)
    role = db.Column(db.String(10), nullable=False)  # admin 或 user

class Category(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(50), nullable=False)
    parent_id = db.Column(db.Integer, db.ForeignKey('category.id'), nullable=True)
    parent = db.relationship('Category', remote_side=[id], backref='children')

class Department(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), unique=True, nullable=False)

class Location(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(100), unique=True, nullable=False)

class AssetStatus(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(50), unique=True, nullable=False)

class Asset(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    asset_number = db.Column(db.String(50), unique=True, nullable=False)  # 资产编号
    name = db.Column(db.String(100), nullable=True)  # 资产名称
    category = db.Column(db.String(50), nullable=False)  # 资产分类
    model = db.Column(db.String(100), nullable=True)  # 型号
    serial_number = db.Column(db.String(100), nullable=True)  # 序列号
    quantity = db.Column(db.Integer, default=1, nullable=False)  # 数量
    department = db.Column(db.String(100), nullable=True)  # 所属部门
    assigned_to = db.Column(db.String(100), nullable=True)  # 使用人
    purchase_date = db.Column(db.Date, nullable=True)  # 购入日期
    warranty_period = db.Column(db.Integer, nullable=True)  # 保修年限
    purchase_amount = db.Column(db.Float, nullable=True)  # 购置金额
    location = db.Column(db.String(100), nullable=True)  # 存放地点
    status = db.Column(db.String(20), nullable=False)  # 资产状态
    remarks = db.Column(db.String(500), nullable=True)  # 备注
    params = db.Column(db.String(500), nullable=True)  # 参数配置

# 资产借用归还模型
class AssetBorrow(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    asset_id = db.Column(db.Integer, db.ForeignKey('asset.id'), nullable=False)  # 资产ID
    borrower = db.Column(db.String(100), nullable=False)  # 借用人
    borrow_date = db.Column(db.Date, nullable=False)  # 借用日期
    return_date = db.Column(db.Date, nullable=True)  # 归还日期
    status = db.Column(db.String(20), nullable=False)  # 状态：借用中/已归还
    remarks = db.Column(db.String(500), nullable=True)  # 备注

# 资产移交模型
class AssetTransfer(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    asset_id = db.Column(db.Integer, db.ForeignKey('asset.id'), nullable=False)  # 资产ID
    from_department = db.Column(db.String(100), nullable=False)  # 转出部门
    to_department = db.Column(db.String(100), nullable=False)  # 转入部门
    from_user = db.Column(db.String(100), nullable=False)  # 转出人
    to_user = db.Column(db.String(100), nullable=False)  # 转入人
    transfer_date = db.Column(db.Date, nullable=False)  # 移交日期
    remarks = db.Column(db.String(500), nullable=True)  # 备注

# 资产操作日志模型
class AssetLog(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    asset_id = db.Column(db.Integer, db.ForeignKey('asset.id'), nullable=False)  # 资产ID
    operation = db.Column(db.String(100), nullable=False)  # 操作类型
    description = db.Column(db.String(500), nullable=False)  # 操作描述
    operation_date = db.Column(db.DateTime, nullable=False, default=datetime.datetime.now)  # 操作日期
    operator = db.Column(db.String(100), nullable=True)  # 操作人

# 生成随机验证码
def generate_captcha():
    # 生成4位随机验证码
    captcha_text = ''.join(random.choices(string.ascii_letters + string.digits, k=4))
    # 将验证码存储在session中
    session['captcha'] = captcha_text
    return captcha_text

# 表单类
class LoginForm(FlaskForm):
    username = StringField('用户名', validators=[InputRequired(), Length(min=4, max=20)])
    password = PasswordField('密码', validators=[InputRequired(), Length(min=6, max=80)])
    captcha = StringField('验证码', validators=[InputRequired(), Length(min=4, max=4)])
    submit = SubmitField('登录')

class CategoryForm(FlaskForm):
    name = StringField('类别名称', validators=[InputRequired(), Length(min=2, max=50)])
    submit = SubmitField('提交')

class AssetForm(FlaskForm):
    asset_number = StringField('资产编号', validators=[InputRequired()])
    name = StringField('资产名称')
    category = StringField('资产分类', validators=[InputRequired()])
    model = StringField('型号')
    serial_number = StringField('序列号')
    quantity = IntegerField('数量', default=1, validators=[InputRequired()])
    department = StringField('所属部门', validators=[InputRequired()])
    assigned_to = StringField('使用人')
    purchase_date = DateField('购入日期', validators=[InputRequired()], format='%Y-%m-%d')
    warranty_period = IntegerField('保修年限', validators=[Optional()])
    purchase_amount = FloatField('购置金额', validators=[Optional()])
    location = StringField('存放地点')
    status = SelectField('资产状态', choices=[('在用', '在用'), ('闲置', '闲置'), ('维修', '维修'), ('报废', '报废')], validators=[InputRequired()])
    remarks = StringField('备注')
    params = TextAreaField('参数配置')
    submit = SubmitField('提交')

# 登录管理器回调
@login_manager.user_loader
def load_user(user_id):
    return User.query.get(int(user_id))

# 路由
@app.route('/')
def index():
    return redirect(url_for('login'))

@app.route('/login', methods=['GET', 'POST'])
def login():
    form = LoginForm()
    # 只在GET请求时生成验证码
    if request.method == 'GET':
        captcha = generate_captcha()
    else:
        # POST请求时使用session中已有的验证码
        captcha = session.get('captcha', '')
    
    if form.validate_on_submit():
        # 验证验证码
        user_captcha = form.captcha.data.lower()
        session_captcha = session.get('captcha', '').lower()
        if user_captcha != session_captcha:
            flash('验证码错误')
            # 重新生成验证码
            captcha = generate_captcha()
            return render_template('login.html', form=form, captcha=captcha)
        # 验证用户名和密码
        user = User.query.filter_by(username=form.username.data).first()
        if user and sha256_crypt.verify(form.password.data, user.password):
            login_user(user)
            return redirect(url_for('dashboard'))
        flash('用户名或密码错误')
        # 重新生成验证码
        captcha = generate_captcha()
    return render_template('login.html', form=form, captcha=captcha)

@app.route('/logout')
@login_required
def logout():
    logout_user()
    return redirect(url_for('login'))

@app.route('/refresh_captcha')
def refresh_captcha():
    # 生成新的验证码
    new_captcha = generate_captcha()
    return {'captcha': new_captcha}

@app.route('/export_assets', methods=['GET', 'POST'])
@login_required
def export_assets():
    # 检查权限
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    
    # 检查请求方法
    if request.method == 'POST':
        # 获取选中的资产ID
        selected_assets_str = request.form.get('selected_assets')
        if not selected_assets_str:
            flash('请先选择要导出的资产')
            return redirect(url_for('dashboard'))
        
        # 解析选中的资产ID
        selected_asset_ids = [int(id) for id in selected_assets_str.split(',')]
        
        # 获取选中的资产
        assets = Asset.query.filter(Asset.id.in_(selected_asset_ids)).all()
        
        # 检查是否有资产可导出
        if not assets:
            flash('没有选中的资产数据可导出')
            return redirect(url_for('dashboard'))
    else:
        # GET请求时，获取所有资产
        assets = Asset.query.all()
        
        # 检查是否有资产可导出
        if not assets:
            flash('没有资产数据可导出')
            return redirect(url_for('dashboard'))
    
    # 生成Excel内容
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from flask import send_file
    import datetime
    import tempfile
    import os
    
    # 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = '资产列表'
    
    # 导入必要的格式模块
    from openpyxl.styles import Font, Alignment, Border, Side
    
    # 定义样式
    # 表头样式：黑体、居中、边框
    header_font = Font(bold=True)
    header_alignment = Alignment(horizontal='center', vertical='center')
    # 内容样式：居中、边框
    content_alignment = Alignment(horizontal='center', vertical='center')
    # 边框样式
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # 写入表头
    headers = ['资产编号*', '资产名称', '资产分类*', '型号', '序列号', '数量', '所属部门', '使用人', '购入日期', '保修年限', '购置金额', '存放地点', '资产状态*', '备注', '参数配置']
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        cell = ws[f'{col_letter}1']
        cell.value = header
        # 应用表头样式
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border
        # 设置列宽为10
        ws.column_dimensions[col_letter].width = 10
    
    # 写入数据
    for row_idx, asset in enumerate(assets, 2):
        # 写入数据
        ws[f'A{row_idx}'] = asset.asset_number
        ws[f'B{row_idx}'] = asset.name or ''
        ws[f'C{row_idx}'] = asset.category
        ws[f'D{row_idx}'] = asset.model or ''
        ws[f'E{row_idx}'] = asset.serial_number or ''
        ws[f'F{row_idx}'] = asset.quantity
        ws[f'G{row_idx}'] = asset.department or ''
        ws[f'H{row_idx}'] = asset.assigned_to or ''
        ws[f'I{row_idx}'] = asset.purchase_date.strftime('%Y-%m-%d') if asset.purchase_date else ''
        ws[f'J{row_idx}'] = asset.warranty_period or ''
        ws[f'K{row_idx}'] = asset.purchase_amount or ''
        ws[f'L{row_idx}'] = asset.location or ''
        ws[f'M{row_idx}'] = asset.status
        ws[f'N{row_idx}'] = asset.remarks or ''
        ws[f'O{row_idx}'] = asset.params or ''
        
        # 应用内容样式
        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            cell = ws[f'{col_letter}{row_idx}']
            cell.alignment = content_alignment
            cell.border = border
        
        # 设置行高为25
        ws.row_dimensions[row_idx].height = 25
    
    # 设置第一行（表头）的行高为25
    ws.row_dimensions[1].height = 25
    
    # 创建临时文件
    with tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False) as f:
        wb.save(f)
        temp_filename = f.name
    
    # 发送文件
    try:
        return send_file(
            temp_filename,
            as_attachment=True,
            download_name=f'资产导出_{datetime.datetime.now().strftime("%Y%m%d%H%M%S")}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    finally:
        # 延迟删除临时文件
        import threading
        def delete_temp_file():
            import time
            time.sleep(10)
            try:
                os.unlink(temp_filename)
            except:
                pass
        
        threading.Thread(target=delete_temp_file).start()

@app.route('/import_assets', methods=['GET', 'POST'])
@login_required
def import_assets():
    # 检查权限
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    
    if request.method == 'POST':
        try:
            # 检查是否有文件上传
            if 'file' not in request.files:
                flash('请选择要导入的文件')
                return redirect(url_for('dashboard'))
            
            file = request.files['file']
            if file.filename == '':
                flash('请选择要导入的文件')
                return redirect(url_for('dashboard'))
            
            # 检查文件扩展名
            if not file.filename.endswith('.xlsx') and not file.filename.endswith('.xls'):
                flash('请上传Excel格式文件（.xlsx或.xls）')
                return redirect(url_for('dashboard'))
            
            # 读取Excel文件
            from openpyxl import load_workbook
            wb = load_workbook(file)
            ws = wb.active
            
            # 验证表头
            expected_headers = ['资产编号', '资产名称', '资产分类', '型号', '序列号', '数量', '所属部门', '使用人', '购入日期', '保修年限', '购置金额', '存放地点', '资产状态', '备注', '参数配置']
            actual_headers = []
            for col in range(1, len(expected_headers) + 1):
                cell_value = ws.cell(row=1, column=col).value
                # 移除表头中的 * 号（如果有的话）
                if cell_value:
                    cell_value = cell_value.replace('*', '')
                actual_headers.append(cell_value)
            
            if actual_headers != expected_headers:
                flash('Excel文件表头格式不正确，请按照导入说明中的格式设置表头')
                return redirect(url_for('dashboard'))
            
            # 读取数据并导入
            import_count = 0
            error_count = 0
            error_messages = []
            
            for row in range(2, ws.max_row + 1):
                try:
                    # 读取单元格数据
                    asset_number = ws.cell(row=row, column=1).value
                    name = ws.cell(row=row, column=2).value
                    category = ws.cell(row=row, column=3).value
                    model = ws.cell(row=row, column=4).value
                    serial_number = ws.cell(row=row, column=5).value
                    quantity = ws.cell(row=row, column=6).value
                    department = ws.cell(row=row, column=7).value
                    assigned_to = ws.cell(row=row, column=8).value
                    purchase_date = ws.cell(row=row, column=9).value
                    warranty_period = ws.cell(row=row, column=10).value
                    purchase_amount = ws.cell(row=row, column=11).value
                    location = ws.cell(row=row, column=12).value
                    status = ws.cell(row=row, column=13).value
                    remarks = ws.cell(row=row, column=14).value
                    params = ws.cell(row=row, column=15).value
                    
                    # 验证必填字段
                    if not asset_number or not category or not status:
                        error_count += 1
                        error_messages.append(f'第{row}行：资产编号、资产分类和资产状态为必填字段')
                        continue
                    
                    # 验证资产编号是否已存在
                    existing_asset = Asset.query.filter_by(asset_number=asset_number).first()
                    if existing_asset:
                        error_count += 1
                        error_messages.append(f'第{row}行：资产编号 {asset_number} 已存在')
                        continue
                    
                    # 验证资产状态
                    valid_statuses = ['在用', '闲置', '维修', '报废']
                    if status not in valid_statuses:
                        error_count += 1
                        error_messages.append(f'第{row}行：资产状态 {status} 无效，必须是：在用、闲置、维修、报废中的一种')
                        continue
                    
                    # 验证资产分类
                    categories = [c.name for c in Category.query.all()]
                    if category not in categories:
                        error_count += 1
                        error_messages.append(f'第{row}行：资产分类 {category} 不存在，请先在系统中添加该分类')
                        continue
                    
                    # 处理日期格式
                    if isinstance(purchase_date, str):
                        try:
                            import datetime
                            purchase_date = datetime.datetime.strptime(purchase_date, '%Y-%m-%d').date()
                        except:
                            error_count += 1
                            error_messages.append(f'第{row}行：购入日期格式不正确，应为 YYYY-MM-DD')
                            continue
                    elif hasattr(purchase_date, 'date'):
                        purchase_date = purchase_date.date()
                    
                    # 创建资产
                    new_asset = Asset(
                        asset_number=asset_number,
                        name=name,
                        category=category,
                        model=model,
                        serial_number=serial_number,
                        quantity=quantity or 1,
                        department=department,
                        assigned_to=assigned_to,
                        purchase_date=purchase_date,
                        warranty_period=warranty_period,
                        purchase_amount=purchase_amount,
                        location=location,
                        status=status,
                        remarks=remarks,
                        params=params
                    )
                    
                    db.session.add(new_asset)
                    # 强制将新资产记录写入数据库，获取资产ID
                    db.session.flush()
                    import_count += 1
                    
                    # 添加操作日志
                    new_log = AssetLog(
                        asset_id=new_asset.id,
                        operation='导入资产',
                        description=f'导入了资产：{new_asset.asset_number} - {new_asset.name}',
                        operator=current_user.username
                    )
                    db.session.add(new_log)
                    
                except Exception as e:
                    error_count += 1
                    error_messages.append(f'第{row}行：导入失败 - {str(e)}')
            
            # 提交事务
            db.session.commit()
            
            # 显示导入结果
            flash(f'资产导入完成：成功 {import_count} 条，失败 {error_count} 条')
            if error_messages:
                # 将所有错误信息合并成一个字符串
                error_message_str = '<div class="error-list"><div class="error-header">错误详情：</div><div class="error-content" style="display: none;"><ul>'
                for error_msg in error_messages:
                    error_message_str += f'<li>{error_msg}</li>'
                error_message_str += '</ul></div><div class="error-toggle" style="cursor: pointer; color: blue; text-decoration: underline;">点击查看详细错误信息</div></div>'
                flash(error_message_str)
            
        except Exception as e:
            flash(f'导入失败：{str(e)}')
    
    return redirect(url_for('dashboard'))

@app.route('/download_import_template')
@login_required
def download_import_template():
    # 检查权限
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    
    # 生成导入模板
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from flask import send_file
    import datetime
    import tempfile
    import os
    
    # 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = '资产导入模板'
    
    # 写入表头
    headers = ['资产编号*', '资产名称', '资产分类*', '型号', '序列号', '数量', '所属部门', '使用人', '购入日期', '保修年限', '购置金额', '存放地点', '资产状态*', '备注', '参数配置']
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        ws[f'{col_letter}1'] = header
    
    # 写入示例数据
    sample_data = [
        ['AST001', '笔记本电脑', '电子设备', 'ThinkPad X1', 'SN123456', 1, '技术部', '张三', '2023-01-01', 3, 8000, '办公室A区', '在用', '办公用', 'i7处理器, 16GB内存, 512GB SSD'],
        ['AST002', '办公椅', '办公家具', '人体工学椅', 'SN654321', 2, '行政部', '', '2023-02-01', 5, 1200, '办公室B区', '闲置', '', '']
    ]
    
    for row_idx, row_data in enumerate(sample_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            col_letter = get_column_letter(col_idx)
            ws[f'{col_letter}{row_idx}'] = value
    
    # 自动调整列宽
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[column].width = adjusted_width
    
    # 创建临时文件
    with tempfile.NamedTemporaryFile(mode='wb', suffix='.xlsx', delete=False) as f:
        wb.save(f)
        temp_filename = f.name
    
    # 发送文件
    try:
        return send_file(
            temp_filename,
            as_attachment=True,
            download_name=f'资产导入模板_{datetime.datetime.now().strftime("%Y%m%d")}.xlsx',
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
    finally:
        # 延迟删除临时文件
        import threading
        def delete_temp_file():
            import time
            time.sleep(10)
            try:
                os.unlink(temp_filename)
            except:
                pass
        
        threading.Thread(target=delete_temp_file).start()

@app.route('/dashboard')
@login_required
def dashboard():
    # 获取过滤参数
    asset_number_filter = request.args.get('asset_number')
    category_filter = request.args.get('category')
    department_filter = request.args.get('department')
    assigned_to_filter = request.args.get('assigned_to')
    content_filter = request.args.get('content')  # 获取要显示的内容区域
    over_5_years_filter = request.args.get('over_5_years')  # 获取超5年资产筛选参数
    over_years_filter = request.args.get('over_years')  # 获取资产年限筛选参数
    status_filter = request.args.get('status')  # 获取状态筛选参数
    page_type = request.args.get('page')  # 获取页面类型参数
    
    # 权限检查：一般用户只能访问资产查询页面
    if current_user.role != 'admin' and (not page_type or page_type != 'query'):
        page_type = 'query'  # 强制设置为查询页面
    
    # 确保查询页面的页面类型保持为query，即使有其他筛选条件
    if request.args.get('page') == 'query':
        page_type = 'query'
    
    # 设置页面标题
    page_title = '资产管理'
    if content_filter == 'category-management':
        page_title = '类别管理'
    elif content_filter == 'department-management':
        page_title = '部门管理'
    elif content_filter == 'user-management':
        page_title = '用户管理'
    elif page_type == 'query':
        page_title = '资产查询'
    elif status_filter:
        if status_filter == '闲置':
            page_title = '闲置资产'
        elif status_filter == '在用':
            page_title = '在用资产'
        elif status_filter == '借用中':
            page_title = '借用中资产'
    elif over_5_years_filter == 'true':
        page_title = '超5年资产'
    
    # 构建查询
    query = Asset.query
    
    # 应用过滤条件
    if asset_number_filter:
        query = query.filter(Asset.asset_number.like('%' + asset_number_filter + '%'))
    if category_filter:
        query = query.filter_by(category=category_filter)
    if department_filter:
        query = query.filter_by(department=department_filter)
    if assigned_to_filter:
        query = query.filter(Asset.assigned_to.like('%' + assigned_to_filter + '%'))
    if status_filter:
        query = query.filter_by(status=status_filter)
    # 资产年限筛选
    if over_years_filter:
        years = int(over_years_filter)
        years_ago = datetime.date.today() - datetime.timedelta(days=years*365.25)
        query = query.filter(Asset.purchase_date < years_ago)
    # 超5年资产筛选（保持兼容性）
    elif over_5_years_filter == 'true':
        five_years_ago = datetime.date.today() - datetime.timedelta(days=5*365.25)
        query = query.filter(Asset.purchase_date < five_years_ago)
    
    # 按资产编号从A-Z，0-9顺序排序
    query = query.order_by(Asset.asset_number)
    
    # 添加分页功能
    page = request.args.get('page_num', 1, type=int)
    per_page = 20  # 每页显示20条数据
    pagination = query.paginate(page=page, per_page=per_page, max_per_page=100, error_out=False)
    assets = pagination.items
    
    # 获取借用中的资产
    borrowed_assets = AssetBorrow.query.filter_by(status='借用中').all()
    
    # 统计不同类型的资产数量（使用数据库查询）
    from sqlalchemy import func
    category_stats_query = db.session.query(Asset.category, func.count(Asset.id)).group_by(Asset.category).all()
    category_stats = {category: count for category, count in category_stats_query}

    # 统计闲置设备（使用数据库查询）
    idle_count = Asset.query.filter_by(status='闲置').count()
    idle_assets = Asset.query.filter_by(status='闲置').limit(100).all()  # 只获取前100个用于显示
    
    # 统计超过5年的设备（使用数据库查询）
    today = datetime.date.today()
    five_years_ago = today - datetime.timedelta(days=5*365.25)
    over_5_years_count = Asset.query.filter(Asset.purchase_date < five_years_ago).count()
    over_5_years_assets = Asset.query.filter(Asset.purchase_date < five_years_ago).limit(100).all()  # 只获取前100个用于显示
    
    # 获取所有类别
    categories = Category.query.all()
    # 获取所有部门
    departments = Department.query.all()
    # 获取所有用户
    users = User.query.all()
    
    # 创建AssetForm实例，用于模态框
    form = AssetForm()
    
    # 获取当前页资产的日志（不再获取所有资产的日志）
    asset_logs = {}
    for asset in assets:
        logs = AssetLog.query.filter_by(asset_id=asset.id).order_by(AssetLog.operation_date.desc()).all()
        # 将 AssetLog 对象转换为可序列化的字典列表
        asset_logs[asset.id] = [{
            'id': log.id,
            'asset_id': log.asset_id,
            'operation': log.operation,
            'description': log.description,
            'operation_date': log.operation_date.strftime('%Y-%m-%d %H:%M:%S') if log.operation_date else '',
            'operator': log.operator
        } for log in logs]
    
    # 为模板功能获取所有资产（包括借用中的资产）
    all_assets = Asset.query.all()
    
    return render_template('dashboard.html', 
                           assets=assets, 
                           all_assets=all_assets,
                           borrowed_assets=borrowed_assets,
                           user=current_user, 
                           category_stats=category_stats,
                           idle_assets=idle_assets,
                           idle_count=idle_count,
                           over_5_years_assets=over_5_years_assets,
                           over_5_years_count=over_5_years_count,
                           asset_logs=asset_logs,
                           asset_number_filter=asset_number_filter,
                           category_filter=category_filter,
                           department_filter=department_filter,
                           assigned_to_filter=assigned_to_filter,
                           content_filter=content_filter,
                           status_filter=status_filter,
                           over_5_years_filter=over_5_years_filter,
                           page_title=page_title,
                           page_type=page_type,
                           form=form,
                           categories=categories,
                           departments=departments,
                           users=users,
                           pagination=pagination,
                           page=page)

# 类别管理功能已集成到dashboard页面的模态框中
# 以下路由函数已被移除，不再需要单独的categories页面

@app.route('/add_asset', methods=['GET', 'POST'])
@login_required
def add_asset():
    if current_user.role != 'admin':
        return '无权限执行此操作'
    form = AssetForm()
    
    # 打印表单数据，用于调试
    print('Form data:', request.form)
    
    if form.validate_on_submit():
        try:
            # 检查资产编号是否已存在
            asset_number = form.asset_number.data
            existing_asset = Asset.query.filter_by(asset_number=asset_number).first()
            if existing_asset:
                return '添加资产失败: 资产编号已存在'
            
            new_asset = Asset(
                asset_number=asset_number,
                name=form.name.data,
                category=form.category.data,
                model=form.model.data,
                serial_number=form.serial_number.data,
                quantity=form.quantity.data,
                department=form.department.data,
                assigned_to=form.assigned_to.data,
                purchase_date=form.purchase_date.data,
                warranty_period=form.warranty_period.data,
                purchase_amount=form.purchase_amount.data,
                location=form.location.data,
                status=form.status.data,
                remarks=form.remarks.data,
                params=form.params.data
            )
            db.session.add(new_asset)
            db.session.commit()
            
            # 添加操作日志
            new_log = AssetLog(
                asset_id=new_asset.id,
                operation='添加资产',
                description=f'添加了资产：{new_asset.asset_number} - {new_asset.name}',
                operator=current_user.username
            )
            db.session.add(new_log)
            db.session.commit()
            
            return '资产添加成功'
        except Exception as e:
            return '添加资产失败: ' + str(e)
    else:
        # 打印验证错误，用于调试
        print('Form errors:', form.errors)
        # 返回详细的验证错误信息
        error_messages = []
        for field, errors in form.errors.items():
            error_messages.append(f'{field}: {" ".join(errors)}')
        return '表单验证失败: ' + '; '.join(error_messages)

@app.route('/add_borrow', methods=['GET', 'POST'])
@login_required
def add_borrow():
    if current_user.role != 'admin':
        return '无权限执行此操作'
    
    if request.method == 'POST':
        try:
            # 获取表单数据
            asset_id = request.form.get('asset_id', type=int)
            borrower = request.form.get('borrower')
            borrow_date = request.form.get('borrow_date')
            return_date = request.form.get('return_date')
            remarks = request.form.get('remarks')
            
            # 验证必填字段
            if not asset_id or not borrower or not borrow_date or not return_date:
                return '表单验证失败: 缺少必填字段'
            
            # 默认状态为借用中
            status = '借用中'
            
            # 检查资产是否存在
            asset = Asset.query.get(asset_id)
            if not asset:
                return '添加借用失败: 资产不存在'
            
            # 如果状态为"借用中"，更新资产状态为"借用中"，并将使用人更新为借用人+"(借用中)"
            if status == '借用中':
                asset.status = '借用中'
                if borrower:
                    asset.assigned_to = borrower + '(借用中)'
                else:
                    asset.assigned_to = '(借用中)'
            
            # 创建借用记录
            new_borrow = AssetBorrow(
                asset_id=asset_id,
                borrower=borrower,
                borrow_date=borrow_date,
                return_date=return_date,
                status=status,
                remarks=remarks
            )
            db.session.add(new_borrow)
            db.session.commit()
            
            # 添加操作日志
            new_log = AssetLog(
                asset_id=asset_id,
                operation='资产借用',
                description=f'借用了资产：{asset.asset_number} - {asset.name or ""}，借用人：{borrower}',
                operator=current_user.username
            )
            db.session.add(new_log)
            db.session.commit()
            
            return '资产借用添加成功'
        except Exception as e:
            return '添加借用失败: ' + str(e)
    
    return '表单验证失败'

@app.route('/add_transfer', methods=['GET', 'POST'])
@login_required
def add_transfer():
    if current_user.role != 'admin':
        return '无权限执行此操作'
    
    if request.method == 'POST':
        try:
            # 获取表单数据
            asset_id = request.form.get('asset_id', type=int)
            from_department = request.form.get('from_department')
            to_department = request.form.get('to_department')
            from_user = request.form.get('from_user')
            to_user = request.form.get('to_user')
            transfer_date = request.form.get('transfer_date')
            remarks = request.form.get('remarks')
            
            # 验证必填字段
            if not asset_id or not to_department or not transfer_date:
                return '表单验证失败: 缺少必填字段'
            
            # 检查资产是否存在
            asset = Asset.query.get(asset_id)
            if not asset:
                return '添加移交失败: 资产不存在'
            
            # 更新资产信息为新的部门和使用人
            asset.department = to_department
            asset.assigned_to = to_user
            
            # 创建移交记录
            new_transfer = AssetTransfer(
                asset_id=asset_id,
                from_department=from_department,
                to_department=to_department,
                from_user=from_user,
                to_user=to_user,
                transfer_date=transfer_date,
                remarks=remarks
            )
            db.session.add(new_transfer)
            db.session.commit()
            
            # 添加操作日志
            new_log = AssetLog(
                asset_id=asset_id,
                operation='资产移交',
                description=f'移交了资产：{asset.asset_number} - {asset.name or ""}，从 {from_department or "未知"} ({from_user or "未知"}) 到 {to_department} ({to_user or "未知"})',
                operator=current_user.username
            )
            db.session.add(new_log)
            db.session.commit()
            
            return '资产移交添加成功'
        except Exception as e:
            return '添加移交失败: ' + str(e)
    
    return '表单验证失败'

@app.route('/return_asset', methods=['POST'])
@login_required
def return_asset():
    if current_user.role != 'admin':
        return '无权限执行此操作'
    
    if request.method == 'POST':
        try:
            # 获取表单数据
            borrow_id = request.form.get('borrow_id')
            
            # 验证必填字段
            if not borrow_id:
                return '表单验证失败: 缺少必填字段'
            
            # 检查借用记录是否存在
            borrow = AssetBorrow.query.get(borrow_id)
            if not borrow:
                return '归还失败: 借用记录不存在'
            
            # 检查资产是否存在
            asset = Asset.query.get(borrow.asset_id)
            if not asset:
                return '归还失败: 资产不存在'
            
            # 更新借用记录状态为已归还
            borrow.status = '已归还'
            borrow.return_date = datetime.date.today()
            
            # 更新资产状态为闲置，使用人设置为待分配
            asset.status = '闲置'
            asset.assigned_to = '待分配'
            
            db.session.commit()
            
            # 添加操作日志
            new_log = AssetLog(
                asset_id=asset.id,
                operation='资产归还',
                description=f'归还了资产：{asset.asset_number} - {asset.name}',
                operator=current_user.username
            )
            db.session.add(new_log)
            db.session.commit()
            
            return '资产归还成功'
        except Exception as e:
            return '归还失败: ' + str(e)
    
    return '表单验证失败'

@app.route('/edit_asset', methods=['POST'])
@login_required
def edit_asset():
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    
    # 从表单中获取资产ID
    asset_id = request.form.get('id')
    if not asset_id:
        flash('资产ID不能为空')
        return redirect(url_for('dashboard'))
    
    asset = Asset.query.get_or_404(asset_id)
    form = AssetForm()
    
    if form.validate_on_submit():
        try:
            # 检查资产编号是否已存在（如果修改了资产编号）
            asset_number = form.asset_number.data
            if asset_number != asset.asset_number:
                existing_asset = Asset.query.filter_by(asset_number=asset_number).first()
                if existing_asset:
                    flash('更新资产失败: 资产编号已存在')
                    return redirect(url_for('dashboard'))
            
            asset.asset_number = asset_number
            asset.name = form.name.data
            asset.category = form.category.data
            asset.model = form.model.data
            asset.serial_number = form.serial_number.data
            asset.quantity = form.quantity.data
            asset.department = form.department.data
            asset.assigned_to = form.assigned_to.data
            asset.purchase_date = form.purchase_date.data
            asset.warranty_period = form.warranty_period.data
            asset.purchase_amount = form.purchase_amount.data
            asset.location = form.location.data
            asset.status = form.status.data
            asset.remarks = form.remarks.data
            asset.params = form.params.data
            db.session.commit()
            
            # 添加操作日志
            new_log = AssetLog(
                asset_id=asset.id,
                operation='编辑资产',
                description=f'编辑了资产：{asset.asset_number} - {asset.name}',
                operator=current_user.username
            )
            db.session.add(new_log)
            db.session.commit()
            
            flash('资产更新成功')
        except Exception as e:
            flash('更新资产失败: ' + str(e))
    else:
        flash('表单验证失败')
    
    return redirect(url_for('dashboard'))

@app.route('/delete_asset/<int:id>')
@login_required
def delete_asset(id):
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    
    try:
        asset = Asset.query.get_or_404(id)
        asset_info = f'{asset.asset_number} - {asset.name}'
        
        # 检查是否有相关的借用记录
        borrows = AssetBorrow.query.filter_by(asset_id=id).all()
        if borrows:
            flash('删除失败：该资产有相关的借用记录，请先处理完借用记录后再删除')
            return redirect(url_for('dashboard'))
        
        # 检查是否有相关的移交记录
        transfers = AssetTransfer.query.filter_by(asset_id=id).all()
        if transfers:
            flash('删除失败：该资产有相关的移交记录，请先处理完移交记录后再删除')
            return redirect(url_for('dashboard'))
        
        # 删除相关的操作日志记录
        AssetLog.query.filter_by(asset_id=id).delete()
        
        db.session.delete(asset)
        db.session.commit()
        
        # 注意：由于资产已经删除，不再添加操作日志
        
        flash('资产删除成功')
    except Exception as e:
        # 回滚事务
        db.session.rollback()
        flash(f'删除资产失败：{str(e)}')
    
    return redirect(url_for('dashboard'))

# 类别管理路由
@app.route('/add_category', methods=['GET', 'POST'])
@login_required
def add_category():
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    form = CategoryForm()
    if request.method == 'POST':
        try:
            # 直接从request获取数据，不使用form.validate_on_submit()
            category_name = request.form.get('name')
            parent_id = request.form.get('parent_id', type=int)
            if category_name:
                # 检查类别名称是否已经存在，特别是在相同父类下
                existing_category = Category.query.filter_by(name=category_name, parent_id=parent_id).first()
                if existing_category:
                    flash('添加类别失败: 类别名称已存在')
                else:
                    new_category = Category(name=category_name, parent_id=parent_id)
                    db.session.add(new_category)
                    db.session.commit()
                    flash('类别添加成功')
            else:
                flash('类别名称不能为空')
        except Exception as e:
            flash('添加类别失败: ' + str(e))
    return redirect(url_for('dashboard', content='category-management'))

@app.route('/edit_category/<int:id>', methods=['GET', 'POST'])
@login_required
def edit_category(id):
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    category = Category.query.get_or_404(id)
    if request.method == 'POST':
        try:
            # 直接从request获取数据
            category_name = request.form.get('name')
            parent_id = request.form.get('parent_id', type=int)
            if category_name:
                category.name = category_name
                category.parent_id = parent_id
                db.session.commit()
                flash('类别更新成功')
            else:
                flash('类别名称不能为空')
        except Exception as e:
            flash('更新类别失败: ' + str(e))
    return redirect(url_for('dashboard', content='category-management'))

@app.route('/delete_category/<int:id>')
@login_required
def delete_category(id):
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    category = Category.query.get_or_404(id)
    
    # 检查是否有子类别
    if category.children:
        flash('删除失败: 该类别下有子类别，请先删除子类别')
        return redirect(url_for('dashboard', content='category-management'))
    
    # 检查是否有资产使用该类别
    if Asset.query.filter_by(category=category.name).first():
        flash('删除失败: 该类别下有资产，请先修改资产类别')
        return redirect(url_for('dashboard', content='category-management'))
    
    db.session.delete(category)
    db.session.commit()
    flash('类别删除成功')
    return redirect(url_for('dashboard', content='category-management'))

# 部门管理路由
@app.route('/add_department', methods=['GET', 'POST'])
@login_required
def add_department():
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        try:
            department_name = request.form.get('name')
            if department_name:
                # 检查部门名称是否已存在
                existing_department = Department.query.filter_by(name=department_name).first()
                if existing_department:
                    flash('添加部门失败: 部门名称已存在')
                else:
                    new_department = Department(name=department_name)
                    db.session.add(new_department)
                    db.session.commit()
                    flash('部门添加成功')
            else:
                flash('部门名称不能为空')
        except Exception as e:
            flash('添加部门失败: ' + str(e))
    return redirect(url_for('dashboard', content='department-management'))

@app.route('/edit_department/<int:id>', methods=['GET', 'POST'])
@login_required
def edit_department(id):
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    department = Department.query.get_or_404(id)
    if request.method == 'POST':
        try:
            department_name = request.form.get('name')
            if department_name:
                # 检查部门名称是否已存在（排除当前部门）
                existing_department = Department.query.filter_by(name=department_name).filter(Department.id != id).first()
                if existing_department:
                    flash('更新部门失败: 部门名称已存在')
                else:
                    department.name = department_name
                    db.session.commit()
                    flash('部门更新成功')
            else:
                flash('部门名称不能为空')
        except Exception as e:
            flash('更新部门失败: ' + str(e))
    return redirect(url_for('dashboard', content='department-management'))

@app.route('/delete_department/<int:id>')
@login_required
def delete_department(id):
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    department = Department.query.get_or_404(id)
    db.session.delete(department)
    db.session.commit()
    flash('部门删除成功')
    return redirect(url_for('dashboard', content='department-management'))

# 用户管理路由
@app.route('/add_user', methods=['GET', 'POST'])
@login_required
def add_user():
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    if request.method == 'POST':
        try:
            # 直接从request获取数据
            username = request.form.get('username')
            password = request.form.get('password')
            role = request.form.get('role')
            if username and password and role:
                # 检查用户名是否已经存在
                existing_user = User.query.filter_by(username=username).first()
                if existing_user:
                    flash('添加用户失败: 用户名已存在')
                else:
                    # 密码加密
                    hashed_password = sha256_crypt.encrypt(password)
                    new_user = User(username=username, password=hashed_password, role=role)
                    db.session.add(new_user)
                    db.session.commit()
                    flash('用户添加成功')
            else:
                flash('用户名、密码和角色不能为空')
        except Exception as e:
            flash('添加用户失败: ' + str(e))
    return redirect(url_for('dashboard', content='user-management'))

@app.route('/edit_user/<int:id>', methods=['GET', 'POST'])
@login_required
def edit_user(id):
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    user = User.query.get_or_404(id)
    if request.method == 'POST':
        try:
            role = request.form.get('role')
            if role:
                user.role = role
                db.session.commit()
                flash('用户更新成功')
            else:
                flash('角色不能为空')
        except Exception as e:
            flash('更新用户失败: ' + str(e))
    return redirect(url_for('dashboard', content='user-management'))

@app.route('/delete_user/<int:id>')
@login_required
def delete_user(id):
    if current_user.role != 'admin':
        flash('无权限执行此操作')
        return redirect(url_for('dashboard'))
    user = User.query.get_or_404(id)
    # 防止删除当前登录的用户
    if user.id == current_user.id:
        flash('不能删除当前登录的用户')
        return redirect(url_for('dashboard', content='user-management'))
    db.session.delete(user)
    db.session.commit()
    flash('用户删除成功')
    return redirect(url_for('dashboard', content='user-management'))

# 初始化数据库
@app.before_first_request
def create_tables():
    # 只创建不存在的表，不删除现有表
    db.create_all()
    
    # 检查默认管理员账号是否存在
    admin = User.query.filter_by(username='admin').first()
    if not admin:
        # 创建默认管理员账号
        admin = User(username='admin', password=sha256_crypt.hash('admin123'), role='admin')
        db.session.add(admin)
    
    # 检查默认普通用户是否存在
    user = User.query.filter_by(username='user').first()
    if not user:
        # 创建默认普通用户
        user = User(username='user', password=sha256_crypt.hash('user123'), role='user')
        db.session.add(user)
    
    # 检查默认资产类别是否存在
    default_categories = ['电子设备', '办公家具', '交通工具', '机械设备', '房屋建筑', '其他资产']
    for category_name in default_categories:
        category = Category.query.filter_by(name=category_name).first()
        if not category:
            category = Category(name=category_name)
            db.session.add(category)
    
    # 提交所有更改
    db.session.commit()

if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)